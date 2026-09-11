"""Celery tasks for report generation.

Handles scheduled and on-demand report generation including PDF, CSV, and Excel formats.
"""

import io
import os
from datetime import datetime, timezone
from typing import Any
from uuid import UUID

import structlog

from app.workers.celery_app import celery_app

logger = structlog.get_logger(__name__)


@celery_app.task(
    name="reports.generate_report",
    bind=True,
    max_retries=2,
    default_retry_delay=30,
    queue="reports",
)
def generate_report(self, report_id: str, report_type: str, params: dict) -> dict:
    """Generate a report based on type and parameters.

    Args:
        report_id: UUID of the report record.
        report_type: Type of report (security_summary, attendance, ppe_compliance, etc.).
        params: Report parameters including date range, cameras, format.

    Returns:
        Dict with status, file_path, and download_url.
    """
    log = logger.bind(report_id=report_id, report_type=report_type)
    log.info("Starting report generation")

    try:
        output_format = params.get("format", "pdf")
        org_id = params.get("org_id", "")
        date_range = {
            "start": params.get("start_date", ""),
            "end": params.get("end_date", ""),
        }

        # Route to appropriate generator
        generators = {
            "security_summary": _generate_security_report,
            "attendance": _generate_attendance_report,
            "ppe_compliance": _generate_ppe_report,
            "vehicle_log": _generate_vehicle_report,
            "footfall": _generate_footfall_report,
        }

        generator_func = generators.get(report_type, _generate_generic_report)
        report_data = generator_func(org_id, date_range, params)

        # Render to requested format
        if output_format == "csv":
            file_bytes = _render_csv(report_data)
            content_type = "text/csv"
            extension = "csv"
        elif output_format == "xlsx":
            file_bytes = _render_excel(report_data)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            extension = "xlsx"
        else:
            file_bytes = _render_pdf(report_data, report_type)
            content_type = "application/pdf"
            extension = "pdf"

        # Upload to MinIO
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        object_name = f"reports/{org_id}/{report_type}_{timestamp}.{extension}"

        try:
            from app.utils.storage import get_storage_client

            storage = get_storage_client()
            storage.upload_bytes(file_bytes, object_name, content_type)
            download_url = storage.get_presigned_url(object_name, expires=86400)
        except Exception as storage_err:
            log.warning("MinIO upload failed, saving locally", error=str(storage_err))
            local_dir = f"/tmp/visionai/reports/{org_id}"
            os.makedirs(local_dir, exist_ok=True)
            local_path = os.path.join(local_dir, f"{report_type}_{timestamp}.{extension}")
            with open(local_path, "wb") as f:
                f.write(file_bytes)
            download_url = local_path
            object_name = local_path

        log.info(
            "Report generated successfully",
            file_size=len(file_bytes),
            format=output_format,
        )

        return {
            "status": "completed",
            "report_id": report_id,
            "file_path": object_name,
            "download_url": download_url,
            "file_size_bytes": len(file_bytes),
            "format": output_format,
        }

    except Exception as exc:
        log.error("Report generation failed", error=str(exc))
        raise self.retry(exc=exc)


def _generate_security_report(org_id: str, date_range: dict, params: dict) -> dict:
    """Generate security summary report data."""
    return {
        "title": "Security Summary Report",
        "type": "security_summary",
        "org_id": org_id,
        "date_range": date_range,
        "sections": [
            {"name": "Overview", "content": "Security overview for the reporting period."},
            {"name": "Alert Summary", "headers": ["Type", "Count", "Critical", "Resolved"],
             "rows": []},
            {"name": "Camera Health", "headers": ["Camera", "Uptime %", "Issues"],
             "rows": []},
            {"name": "Intrusions", "headers": ["Time", "Camera", "Zone", "Status"],
             "rows": []},
        ],
    }


def _generate_attendance_report(org_id: str, date_range: dict, params: dict) -> dict:
    """Generate attendance report data."""
    return {
        "title": "Attendance Report",
        "type": "attendance",
        "org_id": org_id,
        "date_range": date_range,
        "sections": [
            {"name": "Summary", "content": "Attendance summary for the period."},
            {"name": "Daily Attendance",
             "headers": ["Name", "Department", "Date", "First Seen", "Last Seen", "Status"],
             "rows": []},
        ],
    }


def _generate_ppe_report(org_id: str, date_range: dict, params: dict) -> dict:
    """Generate PPE compliance report data."""
    return {
        "title": "PPE Compliance Report",
        "type": "ppe_compliance",
        "org_id": org_id,
        "date_range": date_range,
        "sections": [
            {"name": "Compliance Overview", "content": "PPE compliance summary."},
            {"name": "Violations",
             "headers": ["Time", "Zone", "Type", "Person", "Status"],
             "rows": []},
        ],
    }


def _generate_vehicle_report(org_id: str, date_range: dict, params: dict) -> dict:
    """Generate vehicle log report data."""
    return {
        "title": "Vehicle Activity Report",
        "type": "vehicle_log",
        "org_id": org_id,
        "date_range": date_range,
        "sections": [
            {"name": "Vehicle Summary", "content": "Vehicle activity overview."},
            {"name": "Entry/Exit Log",
             "headers": ["Plate", "Owner", "Entry Time", "Exit Time", "Duration"],
             "rows": []},
        ],
    }


def _generate_footfall_report(org_id: str, date_range: dict, params: dict) -> dict:
    """Generate footfall analytics report data."""
    return {
        "title": "Footfall Analytics Report",
        "type": "footfall",
        "org_id": org_id,
        "date_range": date_range,
        "sections": [
            {"name": "Traffic Summary", "content": "Footfall overview."},
            {"name": "Hourly Data",
             "headers": ["Time", "Zone", "Entries", "Exits", "Occupancy"],
             "rows": []},
        ],
    }


def _generate_generic_report(org_id: str, date_range: dict, params: dict) -> dict:
    """Generate a generic report data structure."""
    return {
        "title": f"Report - {params.get('report_type', 'Custom')}",
        "type": "custom",
        "org_id": org_id,
        "date_range": date_range,
        "sections": [
            {"name": "Report Data", "content": "Custom report data."},
        ],
    }


def _render_pdf(data: dict, report_type: str) -> bytes:
    """Render report data as PDF using reportlab.

    Args:
        data: Report data dictionary.
        report_type: Type identifier.

    Returns:
        PDF file as bytes.
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
        styles = getSampleStyleSheet()
        elements = []

        # Title
        title_style = ParagraphStyle(
            "CustomTitle", parent=styles["Title"], fontSize=20, spaceAfter=12
        )
        elements.append(Paragraph(data.get("title", "Report"), title_style))
        elements.append(Spacer(1, 6 * mm))

        # Date range
        date_range = data.get("date_range", {})
        if date_range.get("start") and date_range.get("end"):
            elements.append(
                Paragraph(
                    f"Period: {date_range['start']} to {date_range['end']}",
                    styles["Normal"],
                )
            )
            elements.append(Spacer(1, 4 * mm))

        # Sections
        for section in data.get("sections", []):
            elements.append(Paragraph(section["name"], styles["Heading2"]))
            elements.append(Spacer(1, 2 * mm))

            if "content" in section:
                elements.append(Paragraph(section["content"], styles["Normal"]))
                elements.append(Spacer(1, 3 * mm))

            if "headers" in section and "rows" in section:
                table_data = [section["headers"]] + section.get("rows", [])
                if len(table_data) > 1:
                    t = Table(table_data, repeatRows=1)
                    t.setStyle(
                        TableStyle(
                            [
                                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.2, 0.3, 0.7)),
                                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                ("FONTSIZE", (0, 0), (-1, 0), 10),
                                ("FONTSIZE", (0, 1), (-1, -1), 9),
                                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.95, 0.95, 0.95)]),
                                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                            ]
                        )
                    )
                    elements.append(t)
                else:
                    elements.append(Paragraph("No data available.", styles["Normal"]))
                elements.append(Spacer(1, 4 * mm))

        # Footer
        elements.append(Spacer(1, 10 * mm))
        footer_style = ParagraphStyle(
            "Footer", parent=styles["Normal"], fontSize=8, textColor=colors.grey
        )
        elements.append(
            Paragraph(
                f"Generated by VisionAI on {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
                footer_style,
            )
        )

        doc.build(elements)
        return buffer.getvalue()

    except ImportError:
        logger.warning("reportlab not available, generating placeholder PDF")
        return b"%PDF-1.4 placeholder - install reportlab for actual PDF generation"


def _render_csv(data: dict) -> bytes:
    """Render report data as CSV.

    Args:
        data: Report data dictionary with sections containing headers and rows.

    Returns:
        CSV file as bytes.
    """
    import csv

    buffer = io.StringIO()
    writer = csv.writer(buffer)

    for section in data.get("sections", []):
        if "headers" in section:
            writer.writerow([f"--- {section['name']} ---"])
            writer.writerow(section["headers"])
            for row in section.get("rows", []):
                writer.writerow(row)
            writer.writerow([])

    return buffer.getvalue().encode("utf-8")


def _render_excel(data: dict) -> bytes:
    """Render report data as Excel (XLSX).

    Args:
        data: Report data dictionary with sections containing headers and rows.

    Returns:
        XLSX file as bytes.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = data.get("title", "Report")[:31]

        header_fill = PatternFill(start_color="334488", end_color="334488", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=10)

        row_num = 1
        for section in data.get("sections", []):
            if "headers" in section:
                # Section header
                ws.cell(row=row_num, column=1, value=section["name"]).font = Font(bold=True, size=12)
                row_num += 1

                # Column headers
                for col, header in enumerate(section["headers"], 1):
                    cell = ws.cell(row=row_num, column=col, value=header)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")
                row_num += 1

                # Data rows
                for row_data in section.get("rows", []):
                    for col, value in enumerate(row_data, 1):
                        ws.cell(row=row_num, column=col, value=value)
                    row_num += 1

                row_num += 1  # Blank row between sections

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    except ImportError:
        logger.warning("openpyxl not available, falling back to CSV")
        return _render_csv(data)


@celery_app.task(
    name="reports.generate_scheduled_reports",
    queue="reports",
)
def generate_scheduled_reports() -> dict:
    """Generate all scheduled reports (called by Celery Beat).

    Returns:
        Dict with count of reports generated.
    """
    log = logger.bind(task="scheduled_reports")
    log.info("Running scheduled report generation")

    # In production, query DB for scheduled report configs and generate each
    generated = 0
    log.info("Scheduled report generation complete", generated=generated)
    return {"generated": generated}
